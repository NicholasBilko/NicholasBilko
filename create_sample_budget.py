from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(__file__).with_name("Simple_Monthly_Budget.xlsx")

NAVY = "17365D"
BLUE = "5B9BD5"
PALE_BLUE = "DDEBF7"
VERY_PALE_BLUE = "F3F8FC"
GREEN = "70AD47"
PALE_GREEN = "E2F0D9"
YELLOW = "FFF2CC"
PALE_RED = "FCE4D6"
RED = "C00000"
ORANGE = "ED7D31"
GRAY = "7F8C8D"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
BLACK = "1F1F1F"

THIN_GRAY = Side(style="thin", color="B7B7B7")
MEDIUM_BLUE = Side(style="medium", color=NAVY)
MONEY = '$#,##0.00;[Red]-$#,##0.00'
MONEY_POSITIVE = '$#,##0.00'
PERCENT = "0%"
MONTH = "mmmm yyyy"
DATE = "mmm d, yyyy"


def table_style() -> TableStyleInfo:
    return TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )


def add_table(ws, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = table_style()
    ws.add_table(table)


def title_cell(cell, text: str, fill: str = NAVY) -> None:
    cell.value = text
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(color=WHITE, bold=True, size=12)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def section_header(ws, cell_range: str, text: str) -> None:
    ws.merge_cells(cell_range)
    first = ws[cell_range.split(":")[0]]
    first.value = text
    first.fill = PatternFill("solid", fgColor=NAVY)
    first.font = Font(color=WHITE, bold=True, size=11)
    first.alignment = Alignment(horizontal="left", vertical="center")


def set_all_borders(ws, cell_range: str) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.border = Border(
                left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
            )


def add_nav_link(ws, cell: str, text: str, target: str) -> None:
    ws[cell] = text
    ws[cell].hyperlink = f"#'{target}'!A1"
    ws[cell].style = "Hyperlink"


def make_start_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Start Here")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 4

    ws.merge_cells("B2:E2")
    ws["B2"] = "Simple Monthly Budget — Sample Workbook"
    ws["B2"].font = Font(size=20, bold=True, color=WHITE)
    ws["B2"].fill = PatternFill("solid", fgColor=NAVY)
    ws["B2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 34

    ws.merge_cells("B3:E3")
    ws["B3"] = (
        "A small month-by-month view that separates spending, reimbursements, "
        "goal contributions, and statement verification."
    )
    ws["B3"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws["B3"].font = Font(color=GRAY, italic=True)
    ws.row_dimensions[3].height = 34

    section_header(ws, "B6:E6", "What the monthly page should answer")
    questions = [
        ("1", "How much income has arrived?"),
        ("2", "What did I personally spend after reimbursements?"),
        ("3", "Am I still within my category budgets?"),
        ("4", "How much did I put toward savings and investments?"),
        ("5", "Do the imported transactions match my bank statements?"),
    ]
    for row, (number, question) in enumerate(questions, start=7):
        ws[f"B{row}"] = number
        ws[f"B{row}"].font = Font(bold=True, color=WHITE)
        ws[f"B{row}"].fill = PatternFill("solid", fgColor=BLUE)
        ws[f"B{row}"].alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        ws.cell(row, 3, question)
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)

    section_header(ws, "B14:E14", "How to try the sample")
    steps = [
        "Open October. Yellow cells are the only cells intended for manual entry.",
        "Review the budget and account-check sections. The statement differences are $0.",
        "Open Review Log and change OCT-014 from Unreviewed to Verified.",
        "Return to October. Excel recalculates the month status to Reconciled.",
        "When adapting this file, keep Power Query output separate and merge reviews by Transaction ID.",
    ]
    for row, step in enumerate(steps, start=15):
        ws[f"B{row}"] = f"{row - 14}."
        ws[f"B{row}"].font = Font(bold=True, color=NAVY)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        ws.cell(row, 3, step)
        ws.cell(row, 3).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 30

    section_header(ws, "B22:E22", "Colour key")
    colour_key = [
        (YELLOW, "Editable input"),
        (PALE_BLUE, "Calculated result"),
        (PALE_GREEN, "Verified or within budget"),
        (PALE_RED, "Needs attention"),
    ]
    for row, (colour, label) in enumerate(colour_key, start=23):
        ws[f"B{row}"].fill = PatternFill("solid", fgColor=colour)
        ws[f"B{row}"].border = Border(
            left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
        )
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        ws.cell(row, 3, label)

    section_header(ws, "B29:E29", "Open a sheet")
    links = [
        ("B30", "October sample", "October"),
        ("C30", "September sample", "September"),
        ("D30", "Transactions", "Transactions"),
        ("E30", "Review Log", "Review Log"),
    ]
    for cell, text, target in links:
        add_nav_link(ws, cell, text, target)
        ws[cell].alignment = Alignment(horizontal="center")

    ws.sheet_properties.tabColor = GREEN
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def make_lists_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Lists")
    statuses = ["Verified", "Unreviewed", "Needs Correction", "Duplicate", "Excluded"]
    types = [
        "Income",
        "Expense",
        "Reimbursement Received",
        "Sinking Fund Contribution",
        "Investment Contribution",
        "Internal Transfer",
    ]
    categories = [
        "Employment",
        "Rent",
        "Groceries",
        "Restaurants",
        "Transportation",
        "Utilities",
        "Subscriptions",
        "Personal",
        "Vacation Fund",
        "TFSA",
        "Transfer",
        "Bank",
    ]
    accounts = ["Chequing", "Savings", "Visa"]

    for col, (header, values) in enumerate(
        [
            ("Statuses", statuses),
            ("Types", types),
            ("Categories", categories),
            ("Accounts", accounts),
        ],
        start=1,
    ):
        ws.cell(1, col, header)
        for row, value in enumerate(values, start=2):
            ws.cell(row, col, value)
    wb.defined_names.add(
        DefinedName("StatusList", attr_text="'Lists'!$A$2:$A$6")
    )
    wb.defined_names.add(
        DefinedName("TypeList", attr_text="'Lists'!$B$2:$B$7")
    )
    wb.defined_names.add(
        DefinedName("CategoryList", attr_text="'Lists'!$C$2:$C$13")
    )
    ws.sheet_state = "hidden"


def sample_data():
    transactions = [
        # September
        ("SEP-001", date(2026, 9, 1), "Payroll Deposit", "Chequing", 4000.00, "Income", "Employment"),
        ("SEP-002", date(2026, 9, 2), "E-transfer to Landlord", "Chequing", -2000.00, "Expense", "Rent"),
        ("SEP-003", date(2026, 9, 2), "E-transfer from Roommate", "Chequing", 1000.00, "Internal Transfer", "Bank"),
        ("SEP-004", date(2026, 9, 3), "Transfer to Vacation Savings", "Chequing", -200.00, "Internal Transfer", "Transfer"),
        ("SEP-005", date(2026, 9, 3), "Transfer from Chequing", "Savings", 200.00, "Internal Transfer", "Transfer"),
        ("SEP-006", date(2026, 9, 4), "Investment Account Contribution", "Chequing", -400.00, "Internal Transfer", "Transfer"),
        ("SEP-007", date(2026, 9, 5), "Neighbourhood Grocery", "Visa", -210.00, "Expense", "Groceries"),
        ("SEP-008", date(2026, 9, 6), "Fuel Station", "Visa", -60.00, "Expense", "Transportation"),
        ("SEP-009", date(2026, 9, 8), "Hydro Company", "Chequing", -135.00, "Expense", "Utilities"),
        ("SEP-010", date(2026, 9, 10), "Local Restaurant", "Visa", -80.00, "Expense", "Restaurants"),
        ("SEP-011", date(2026, 9, 12), "Streaming Service", "Visa", -15.00, "Expense", "Subscriptions"),
        ("SEP-012", date(2026, 9, 15), "Visa Payment", "Chequing", -500.00, "Internal Transfer", "Transfer"),
        ("SEP-013", date(2026, 9, 15), "Payment Received", "Visa", 500.00, "Internal Transfer", "Transfer"),
        ("SEP-014", date(2026, 9, 18), "Neighbourhood Grocery", "Visa", -160.00, "Expense", "Groceries"),
        ("SEP-015", date(2026, 9, 20), "Personal Purchase", "Visa", -90.00, "Expense", "Personal"),
        ("SEP-016", date(2026, 9, 25), "Local Restaurant", "Visa", -120.00, "Expense", "Restaurants"),
        ("SEP-017", date(2026, 9, 26), "E-transfer from Friend", "Chequing", 60.00, "Internal Transfer", "Bank"),
        ("SEP-018", date(2026, 9, 28), "Neighbourhood Grocery", "Visa", -140.00, "Expense", "Groceries"),
        ("SEP-019", date(2026, 9, 30), "Internet Provider", "Chequing", -70.00, "Expense", "Utilities"),
        # October
        ("OCT-001", date(2026, 10, 1), "Payroll Deposit", "Chequing", 4000.00, "Income", "Employment"),
        ("OCT-002", date(2026, 10, 2), "E-transfer to Landlord", "Chequing", -2000.00, "Expense", "Rent"),
        ("OCT-003", date(2026, 10, 2), "E-transfer from Roommate", "Chequing", 1000.00, "Internal Transfer", "Bank"),
        ("OCT-004", date(2026, 10, 3), "Neighbourhood Grocery", "Visa", -185.40, "Expense", "Groceries"),
        ("OCT-005", date(2026, 10, 4), "Dinner with Friends", "Visa", -150.00, "Expense", "Restaurants"),
        ("OCT-006", date(2026, 10, 5), "E-transfer from Friend", "Chequing", 100.00, "Internal Transfer", "Bank"),
        ("OCT-007", date(2026, 10, 5), "Transfer to Vacation Savings", "Chequing", -300.00, "Internal Transfer", "Transfer"),
        ("OCT-008", date(2026, 10, 5), "Transfer from Chequing", "Savings", 300.00, "Internal Transfer", "Transfer"),
        ("OCT-009", date(2026, 10, 6), "Investment Account Contribution", "Chequing", -500.00, "Internal Transfer", "Transfer"),
        ("OCT-010", date(2026, 10, 7), "Fuel Station", "Visa", -72.00, "Expense", "Transportation"),
        ("OCT-011", date(2026, 10, 8), "Hydro Company", "Chequing", -140.00, "Expense", "Utilities"),
        ("OCT-012", date(2026, 10, 9), "Streaming Service", "Visa", -15.00, "Expense", "Subscriptions"),
        ("OCT-013", date(2026, 10, 10), "Visa Payment", "Chequing", -500.00, "Internal Transfer", "Transfer"),
        ("OCT-013B", date(2026, 10, 10), "Payment Received", "Visa", 500.00, "Internal Transfer", "Transfer"),
        ("OCT-014", date(2026, 10, 12), "Personal Purchase", "Visa", -65.00, "Expense", "Personal"),
        ("OCT-015", date(2026, 10, 13), "Neighbourhood Grocery", "Visa", -95.60, "Expense", "Groceries"),
    ]

    review_rows = []
    for transaction in transactions:
        transaction_id = transaction[0]
        posted_date = transaction[1]
        review_rows.append(
            [
                transaction_id,
                "Unreviewed" if transaction_id == "OCT-014" else "Verified",
                "",
                "",
                "",
                date(posted_date.year, posted_date.month, 1),
                "",
            ]
        )

    overrides = {
        "SEP-002": ("", "", "RENT-2026-09", "Full rent paid"),
        "SEP-003": ("Reimbursement Received", "Rent", "RENT-2026-09", "Roommate's share"),
        "SEP-004": ("Sinking Fund Contribution", "Vacation Fund", "VACATION-2026-09", ""),
        "SEP-005": ("", "", "VACATION-2026-09", "Matching account transfer"),
        "SEP-006": ("Investment Contribution", "TFSA", "INVEST-2026-09", ""),
        "SEP-012": ("", "", "VISA-PAYMENT-2026-09", ""),
        "SEP-013": ("", "", "VISA-PAYMENT-2026-09", ""),
        "SEP-016": ("", "", "DINNER-2026-09-25", "Dinner paid in full"),
        "SEP-017": ("Reimbursement Received", "Restaurants", "DINNER-2026-09-25", "Friend's share"),
        "OCT-002": ("", "", "RENT-2026-10", "Full rent paid"),
        "OCT-003": ("Reimbursement Received", "Rent", "RENT-2026-10", "Roommate's share"),
        "OCT-005": ("", "", "DINNER-2026-10-04", "Dinner paid in full"),
        "OCT-006": ("Reimbursement Received", "Restaurants", "DINNER-2026-10-04", "Friends' shares"),
        "OCT-007": ("Sinking Fund Contribution", "Vacation Fund", "VACATION-2026-10", ""),
        "OCT-008": ("", "", "VACATION-2026-10", "Matching account transfer"),
        "OCT-009": ("Investment Contribution", "TFSA", "INVEST-2026-10", ""),
        "OCT-013": ("", "", "VISA-PAYMENT-2026-10", ""),
        "OCT-013B": ("", "", "VISA-PAYMENT-2026-10", ""),
    }
    review_by_id = {row[0]: row for row in review_rows}
    for transaction_id, (type_override, category_override, related, notes) in overrides.items():
        row = review_by_id[transaction_id]
        row[2] = type_override
        row[3] = category_override
        row[4] = related
        row[6] = notes

    return transactions, review_rows


def make_review_sheet(wb: Workbook, review_rows) -> None:
    ws = wb.create_sheet("Review Log")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:G1")
    ws["A1"] = "Transaction Review & Overrides"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:G2")
    ws["A2"] = (
        "Edit this yellow table—not imported Power Query rows. "
        "Transaction ID keeps every decision attached after refresh."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=YELLOW)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 32
    add_nav_link(ws, "A3", "← Start Here", "Start Here")
    add_nav_link(ws, "B3", "October", "October")
    add_nav_link(ws, "C3", "September", "September")
    add_nav_link(ws, "D3", "Transactions", "Transactions")

    headers = [
        "Transaction ID",
        "Status",
        "Type Override",
        "Category Override",
        "Related Transaction",
        "Applies To Month",
        "Notes",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(4, col, header)
    for row_number, values in enumerate(review_rows, start=5):
        for col, value in enumerate(values, start=1):
            ws.cell(row_number, col, value)
            ws.cell(row_number, col).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row_number, 6).number_format = MONTH

    last_row = 4 + len(review_rows)
    add_table(ws, "tblReview", f"A4:G{last_row}")
    widths = [16, 20, 30, 24, 26, 20, 36]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col)].width = width

    status_validation = DataValidation(
        type="list", formula1="=StatusList", allow_blank=False
    )
    type_validation = DataValidation(
        type="list", formula1="=TypeList", allow_blank=True
    )
    category_validation = DataValidation(
        type="list", formula1="=CategoryList", allow_blank=True
    )
    for validation in [status_validation, type_validation, category_validation]:
        ws.add_data_validation(validation)
    status_validation.add(f"B5:B{last_row}")
    type_validation.add(f"C5:C{last_row}")
    category_validation.add(f"D5:D{last_row}")

    ws.conditional_formatting.add(
        f"B5:B{last_row}",
        FormulaRule(
            formula=['B5="Verified"'],
            fill=PatternFill("solid", fgColor=PALE_GREEN),
        ),
    )
    ws.conditional_formatting.add(
        f"B5:B{last_row}",
        FormulaRule(
            formula=['OR(B5="Unreviewed",B5="Needs Correction")'],
            fill=PatternFill("solid", fgColor=PALE_RED),
        ),
    )
    ws.auto_filter.ref = f"A4:G{last_row}"
    ws.sheet_properties.tabColor = ORANGE


def make_transactions_sheet(wb: Workbook, transactions) -> None:
    ws = wb.create_sheet("Transactions")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:M1")
    ws["A1"] = "Final Transactions"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:M2")
    ws["A2"] = (
        "Columns A:G represent cleaned bank data. Columns H:M apply the Review Log. "
        "Monthly pages use only the final columns."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 32
    add_nav_link(ws, "A3", "← Start Here", "Start Here")
    add_nav_link(ws, "B3", "October", "October")
    add_nav_link(ws, "C3", "September", "September")
    add_nav_link(ws, "D3", "Review Log", "Review Log")

    headers = [
        "Transaction ID",
        "Posted Date",
        "Merchant",
        "Account",
        "Amount",
        "Suggested Type",
        "Suggested Category",
        "Final Type",
        "Final Category",
        "Related Transaction",
        "Applies To Month",
        "Review Status",
        "Notes",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(4, col, header)

    for row_number, transaction in enumerate(transactions, start=5):
        for col, value in enumerate(transaction, start=1):
            ws.cell(row_number, col, value)
        ws.cell(row_number, 2).number_format = DATE
        ws.cell(row_number, 5).number_format = MONEY
        ws.cell(
            row_number,
            8,
            f'=IFERROR(IF(INDEX(tblReview[Type Override],MATCH(A{row_number},'
            f'tblReview[Transaction ID],0))="",F{row_number},'
            f'INDEX(tblReview[Type Override],MATCH(A{row_number},'
            f'tblReview[Transaction ID],0))),F{row_number})',
        )
        ws.cell(
            row_number,
            9,
            f'=IFERROR(IF(INDEX(tblReview[Category Override],MATCH(A{row_number},'
            f'tblReview[Transaction ID],0))="",G{row_number},'
            f'INDEX(tblReview[Category Override],MATCH(A{row_number},'
            f'tblReview[Transaction ID],0))),G{row_number})',
        )
        ws.cell(
            row_number,
            10,
            f'=IFERROR(INDEX(tblReview[Related Transaction],MATCH(A{row_number},'
            f'tblReview[Transaction ID],0)),"")',
        )
        ws.cell(
            row_number,
            11,
            f'=IFERROR(IF(INDEX(tblReview[Applies To Month],MATCH(A{row_number},'
            f'tblReview[Transaction ID],0))=0,DATE(YEAR(B{row_number}),'
            f'MONTH(B{row_number}),1),INDEX(tblReview[Applies To Month],'
            f'MATCH(A{row_number},tblReview[Transaction ID],0))),'
            f'DATE(YEAR(B{row_number}),MONTH(B{row_number}),1))',
        )
        ws.cell(row_number, 11).number_format = MONTH
        ws.cell(
            row_number,
            12,
            f'=IFERROR(INDEX(tblReview[Status],MATCH(A{row_number},'
            f'tblReview[Transaction ID],0)),"Unreviewed")',
        )
        ws.cell(
            row_number,
            13,
            f'=IFERROR(INDEX(tblReview[Notes],MATCH(A{row_number},'
            f'tblReview[Transaction ID],0)),"")',
        )
        for col in range(8, 14):
            ws.cell(row_number, col).fill = PatternFill("solid", fgColor=PALE_BLUE)

    last_row = 4 + len(transactions)
    add_table(ws, "tblTransactions", f"A4:M{last_row}")
    widths = [16, 15, 32, 16, 15, 28, 24, 30, 24, 26, 20, 20, 34]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col)].width = width
    ws.conditional_formatting.add(
        f"L5:L{last_row}",
        FormulaRule(
            formula=['L5="Verified"'],
            fill=PatternFill("solid", fgColor=PALE_GREEN),
        ),
    )
    ws.conditional_formatting.add(
        f"L5:L{last_row}",
        FormulaRule(
            formula=['L5<>"Verified"'],
            fill=PatternFill("solid", fgColor=PALE_RED),
        ),
    )
    ws.sheet_properties.tabColor = GRAY


def make_month_sheet(
    wb: Workbook,
    sheet_name: str,
    month_start: date,
    budgets: dict[str, float],
    expected_income: float,
    sinking_goal: float,
    investment_goal: float,
    opening_balances: dict[str, float],
    statement_balances: dict[str, float],
) -> None:
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A10"
    ws.sheet_properties.tabColor = BLUE
    ws["L1"] = month_start
    ws["L1"].number_format = MONTH
    ws.column_dimensions["L"].hidden = True

    widths = {
        "A": 22,
        "B": 16,
        "C": 17,
        "D": 17,
        "E": 17,
        "F": 17,
        "G": 13,
        "H": 3,
        "I": 23,
        "J": 20,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:J1")
    ws["A1"] = f"{sheet_name} 2026 — Monthly Budget Check"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells("A2:J2")
    ws["A2"] = (
        "A simple answer to: what came in, what I spent, what I set aside, "
        "what remains, and whether the month matches the bank."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
    ws["A2"].font = Font(color=GRAY, italic=True)
    ws["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[2].height = 30

    add_nav_link(ws, "A3", "← Start Here", "Start Here")
    other_month = "September" if sheet_name == "October" else "October"
    add_nav_link(ws, "B3", other_month, other_month)
    add_nav_link(ws, "C3", "Transactions", "Transactions")
    add_nav_link(ws, "D3", "Review Log", "Review Log")
    ws.merge_cells("F3:J3")
    ws["F3"] = "Yellow = editable   |   Blue = calculated"
    ws["F3"].font = Font(italic=True, color=GRAY)
    ws["F3"].alignment = Alignment(horizontal="right")

    cards = [
        ("A5:B5", "A6:B7", "A8:B8", "Income Received"),
        ("C5:D5", "C6:D7", "C8:D8", "Net Spending"),
        ("E5:F5", "E6:F7", "E8:F8", "Budget Remaining"),
        ("G5:H5", "G6:H7", "G8:H8", "Goals Funded"),
        ("I5:J5", "I6:J7", "I8:J8", "Unallocated Income"),
    ]
    for header_range, value_range, detail_range, label in cards:
        ws.merge_cells(header_range)
        ws.merge_cells(value_range)
        ws.merge_cells(detail_range)
        header = ws[header_range.split(":")[0]]
        value = ws[value_range.split(":")[0]]
        detail = ws[detail_range.split(":")[0]]
        header.value = label
        header.fill = PatternFill("solid", fgColor=BLUE)
        header.font = Font(color=WHITE, bold=True)
        header.alignment = Alignment(horizontal="center")
        value.fill = PatternFill("solid", fgColor=PALE_BLUE)
        value.font = Font(color=NAVY, bold=True, size=18)
        value.alignment = Alignment(horizontal="center", vertical="center")
        value.number_format = MONEY_POSITIVE
        detail.fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
        detail.font = Font(color=GRAY, italic=True, size=9)
        detail.alignment = Alignment(horizontal="center")
        set_all_borders(ws, f"{header_range.split(':')[0]}:{detail_range.split(':')[1]}")

    income_formula = (
        '=SUMIFS(tblTransactions[Amount],tblTransactions[Final Type],"Income",'
        'tblTransactions[Applies To Month],$L$1)'
    )
    gross_formula = (
        '=-SUMIFS(tblTransactions[Amount],tblTransactions[Final Type],"Expense",'
        'tblTransactions[Applies To Month],$L$1)'
    )
    reimb_formula = (
        '=SUMIFS(tblTransactions[Amount],tblTransactions[Final Type],'
        '"Reimbursement Received",tblTransactions[Applies To Month],$L$1)'
    )
    sinking_formula = (
        '=-SUMIFS(tblTransactions[Amount],tblTransactions[Final Type],'
        '"Sinking Fund Contribution",tblTransactions[Applies To Month],$L$1)'
    )
    investment_formula = (
        '=-SUMIFS(tblTransactions[Amount],tblTransactions[Final Type],'
        '"Investment Contribution",tblTransactions[Applies To Month],$L$1)'
    )

    ws["A6"] = "=J23"
    ws["C6"] = "=J26"
    ws["E6"] = "=SUM(F12:F18)"
    ws["G6"] = "=J27+J28"
    ws["I6"] = "=J29"
    ws["A8"] = '="of "&TEXT(J11,"$#,##0")&" expected"'
    ws["C8"] = '="Gross "&TEXT(SUM(C12:C18),"$#,##0")&" less "&TEXT(SUM(D12:D18),"$#,##0")'
    ws["E8"] = '="of "&TEXT(SUM(B12:B18),"$#,##0")&" category budget"'
    ws["G8"] = '="Goal "&TEXT(J12+J13,"$#,##0")'
    ws["I8"] = "After spending and goal contributions"

    section_header(ws, "A10:G10", "Budget by category — edit only the yellow Budget cells")
    budget_headers = [
        "Category",
        "Budget",
        "Gross Spend",
        "Reimbursed",
        "Net Spend",
        "Remaining",
        "% Used",
    ]
    for col, header in enumerate(budget_headers, start=1):
        ws.cell(11, col, header)
        ws.cell(11, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(11, col).font = Font(color=WHITE, bold=True)
        ws.cell(11, col).alignment = Alignment(horizontal="center", wrap_text=True)

    for row, (category, budget) in enumerate(budgets.items(), start=12):
        ws.cell(row, 1, category)
        ws.cell(row, 2, budget)
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(
            row,
            3,
            f'=-SUMIFS(tblTransactions[Amount],tblTransactions[Final Type],"Expense",'
            f'tblTransactions[Final Category],$A{row},'
            f'tblTransactions[Applies To Month],$L$1)',
        )
        ws.cell(
            row,
            4,
            f'=SUMIFS(tblTransactions[Amount],tblTransactions[Final Type],'
            f'"Reimbursement Received",tblTransactions[Final Category],$A{row},'
            f'tblTransactions[Applies To Month],$L$1)',
        )
        ws.cell(row, 5, f"=C{row}-D{row}")
        ws.cell(row, 6, f"=B{row}-E{row}")
        ws.cell(row, 7, f'=IFERROR(E{row}/B{row},0)')
        for col in range(2, 7):
            ws.cell(row, col).number_format = MONEY_POSITIVE
        ws.cell(row, 7).number_format = PERCENT
        for col in range(3, 8):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)

    total_row = 12 + len(budgets)
    ws.cell(total_row, 1, "TOTAL")
    ws.cell(total_row, 1).font = Font(bold=True)
    for col in range(2, 7):
        letter = chr(64 + col)
        ws.cell(total_row, col, f"=SUM({letter}12:{letter}{total_row - 1})")
        ws.cell(total_row, col).number_format = MONEY_POSITIVE
        ws.cell(total_row, col).font = Font(bold=True)
        ws.cell(total_row, col).fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws.cell(total_row, 7, f'=IFERROR(E{total_row}/B{total_row},0)')
    ws.cell(total_row, 7).number_format = PERCENT
    ws.cell(total_row, 7).font = Font(bold=True)
    ws.cell(total_row, 7).fill = PatternFill("solid", fgColor=PALE_BLUE)
    set_all_borders(ws, f"A11:G{total_row}")

    ws.conditional_formatting.add(
        f"F12:F{total_row}",
        CellIsRule(
            operator="lessThan",
            formula=["0"],
            fill=PatternFill("solid", fgColor=PALE_RED),
            font=Font(color=RED),
        ),
    )
    ws.conditional_formatting.add(
        f"F12:F{total_row}",
        CellIsRule(
            operator="greaterThanOrEqual",
            formula=["0"],
            fill=PatternFill("solid", fgColor=PALE_GREEN),
        ),
    )
    ws.conditional_formatting.add(
        f"G12:G{total_row}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=BLUE),
    )

    section_header(ws, "I10:J10", "Month setup & verification")
    setup_labels = [
        "Expected income",
        "Sinking fund goal",
        "Investment goal",
        "Transactions through",
        "Transactions",
        "Verified",
        "Needs review",
        "Statement difference",
        "Month status",
    ]
    for row, label in enumerate(setup_labels, start=11):
        ws.cell(row, 9, label)
        ws.cell(row, 9).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
        ws.cell(row, 9).font = Font(bold=row == 19)
        ws.cell(row, 10).fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws.cell(row, 10).alignment = Alignment(horizontal="center")
    ws["J11"] = expected_income
    ws["J12"] = sinking_goal
    ws["J13"] = investment_goal
    for cell in ["J11", "J12", "J13"]:
        ws[cell].fill = PatternFill("solid", fgColor=YELLOW)
        ws[cell].number_format = MONEY_POSITIVE
        ws[cell].comment = Comment("Editable planning input.", "Cursor")
    ws["J14"] = (
        '=IFERROR(AGGREGATE(14,6,tblTransactions[Posted Date]/'
        '(tblTransactions[Applies To Month]=$L$1),1),"")'
    )
    ws["J14"].number_format = DATE
    ws["J15"] = (
        '=COUNTIFS(tblTransactions[Applies To Month],$L$1,'
        'tblTransactions[Review Status],"<>Excluded")'
    )
    ws["J16"] = (
        '=COUNTIFS(tblTransactions[Applies To Month],$L$1,'
        'tblTransactions[Review Status],"Verified")'
    )
    ws["J17"] = "=J15-J16"
    ws["J18"] = "=SUMPRODUCT(ABS(F24:F26))"
    ws["J18"].number_format = MONEY_POSITIVE
    ws["J19"] = '=IF(AND(J17=0,J18<0.01),"RECONCILED","NEEDS REVIEW")'
    ws["J19"].font = Font(bold=True)
    set_all_borders(ws, "I10:J19")
    ws.conditional_formatting.add(
        "J19",
        FormulaRule(
            formula=['J19="RECONCILED"'],
            fill=PatternFill("solid", fgColor=PALE_GREEN),
            font=Font(color=GREEN, bold=True),
        ),
    )
    ws.conditional_formatting.add(
        "J19",
        FormulaRule(
            formula=['J19<>"RECONCILED"'],
            fill=PatternFill("solid", fgColor=PALE_RED),
            font=Font(color=RED, bold=True),
        ),
    )

    section_header(ws, "A22:G22", "Bank statement check — edit yellow opening and statement balances")
    reconciliation_headers = [
        "Account",
        "Opening Balance",
        "Imported Activity",
        "Calculated Ending",
        "Statement Ending",
        "Difference",
        "Status",
    ]
    for col, header in enumerate(reconciliation_headers, start=1):
        ws.cell(23, col, header)
        ws.cell(23, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(23, col).font = Font(color=WHITE, bold=True)
        ws.cell(23, col).alignment = Alignment(horizontal="center", wrap_text=True)

    accounts = ["Chequing", "Savings", "Visa"]
    for row, account in enumerate(accounts, start=24):
        ws.cell(row, 1, account)
        ws.cell(row, 2, opening_balances[account])
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(
            row,
            3,
            f'=SUMIFS(tblTransactions[Amount],tblTransactions[Account],$A{row},'
            f'tblTransactions[Posted Date],">="&$L$1,'
            f'tblTransactions[Posted Date],"<"&EDATE($L$1,1))',
        )
        ws.cell(row, 4, f"=B{row}+C{row}")
        ws.cell(row, 5, statement_balances[account])
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row, 6, f"=E{row}-D{row}")
        ws.cell(row, 7, f'=IF(ABS(F{row})<0.01,"Matched","Check")')
        for col in range(2, 7):
            ws.cell(row, col).number_format = MONEY
        for col in [3, 4, 6, 7]:
            ws.cell(row, col).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
    set_all_borders(ws, "A23:G26")
    ws.conditional_formatting.add(
        "F24:F26",
        FormulaRule(
            formula=["ABS(F24)<0.01"],
            fill=PatternFill("solid", fgColor=PALE_GREEN),
        ),
    )
    ws.conditional_formatting.add(
        "F24:F26",
        FormulaRule(
            formula=["ABS(F24)>=0.01"],
            fill=PatternFill("solid", fgColor=PALE_RED),
        ),
    )

    section_header(ws, "I22:J22", "How the month adds up")
    allocation_labels = [
        "Income received",
        "Gross spending",
        "Reimbursements",
        "Net spending",
        "Sinking funds",
        "Investments",
        "Unallocated income",
    ]
    allocation_formulas = [
        income_formula,
        gross_formula,
        reimb_formula,
        "=J24-J25",
        sinking_formula,
        investment_formula,
        "=J23-J26-J27-J28",
    ]
    for row, (label, formula) in enumerate(
        zip(allocation_labels, allocation_formulas), start=23
    ):
        ws.cell(row, 9, label)
        ws.cell(row, 10, formula)
        ws.cell(row, 10).number_format = MONEY_POSITIVE
        ws.cell(row, 9).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
        ws.cell(row, 10).fill = PatternFill(
            "solid", fgColor=PALE_GREEN if row == 29 else PALE_BLUE
        )
        if row == 29:
            ws.cell(row, 9).font = Font(bold=True)
            ws.cell(row, 10).font = Font(bold=True)
    set_all_borders(ws, "I22:J29")

    ws.merge_cells("A31:J31")
    ws["A31"] = (
        "Reminder: contributions are allocations, not spending. "
        "Internal transfers are included for bank reconciliation but excluded from the budget."
    )
    ws["A31"].fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    ws["A31"].font = Font(color=GRAY, italic=True)
    ws["A31"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[31].height = 28

    for row in range(1, 32):
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 20)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:J31"
    ws.auto_filter.ref = f"A11:G{total_row}"


def build_workbook() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    make_start_sheet(wb)
    make_lists_sheet(wb)
    transactions, review_rows = sample_data()
    make_review_sheet(wb, review_rows)
    make_transactions_sheet(wb, transactions)

    common_budgets = {
        "Rent": 1000.00,
        "Groceries": 500.00,
        "Restaurants": 200.00,
        "Transportation": 180.00,
        "Utilities": 220.00,
        "Subscriptions": 50.00,
        "Personal": 250.00,
    }
    september_budgets = common_budgets.copy()
    september_budgets["Restaurants"] = 180.00

    make_month_sheet(
        wb,
        "September",
        date(2026, 9, 1),
        september_budgets,
        expected_income=4000.00,
        sinking_goal=200.00,
        investment_goal=400.00,
        opening_balances={"Chequing": 2500.00, "Savings": 1800.00, "Visa": -250.00},
        statement_balances={"Chequing": 4255.00, "Savings": 2000.00, "Visa": -625.00},
    )
    make_month_sheet(
        wb,
        "October",
        date(2026, 10, 1),
        common_budgets,
        expected_income=4000.00,
        sinking_goal=300.00,
        investment_goal=500.00,
        opening_balances={"Chequing": 4255.00, "Savings": 2000.00, "Visa": -625.00},
        statement_balances={"Chequing": 5915.00, "Savings": 2300.00, "Visa": -708.00},
    )

    # Put the useful sheets first while keeping the supporting data nearby.
    desired_order = [
        "Start Here",
        "October",
        "September",
        "Transactions",
        "Review Log",
        "Lists",
    ]
    wb._sheets = [wb[name] for name in desired_order]
    wb.active = wb.sheetnames.index("October")
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    expected_sheets = [
        "Start Here",
        "October",
        "September",
        "Transactions",
        "Review Log",
        "Lists",
    ]
    assert wb.sheetnames == expected_sheets
    assert wb["Lists"].sheet_state == "hidden"
    assert "tblTransactions" in wb["Transactions"].tables
    assert "tblReview" in wb["Review Log"].tables
    assert wb["October"]["J19"].value.startswith("=IF(")
    assert wb["October"]["A6"].value == "=J23"
    assert wb["October"]["B12"].fill.fgColor.rgb.endswith(YELLOW)
    assert wb["October"]["F24"].value == "=E24-D24"
    assert wb["Review Log"]["B38"].value == "Unreviewed"


if __name__ == "__main__":
    build_workbook()
    validate_workbook()
    print(f"Created and validated {OUTPUT}")
