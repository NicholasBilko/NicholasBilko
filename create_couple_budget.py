from __future__ import annotations

from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from create_sample_budget import (
    BLUE,
    DATE,
    GRAY,
    GREEN,
    LIGHT_GRAY,
    MONEY,
    MONEY_POSITIVE,
    MONTH,
    NAVY,
    ORANGE,
    PALE_BLUE,
    PALE_GREEN,
    PALE_RED,
    PERCENT,
    RED,
    THIN_GRAY,
    VERY_PALE_BLUE,
    WHITE,
    YELLOW,
    add_nav_link,
    add_table,
    section_header,
    set_all_borders,
)


OUTPUT = Path(__file__).with_name("Couple_Monthly_Budget.xlsx")


ACCOUNTS = [
    ("CIBC-P1-CHQ", "CIBC", "Everyday Chequing", "Partner 1", "Chequing", "Yes", "Yes"),
    ("RBC-P2-CHQ", "RBC", "Day-to-Day Chequing", "Partner 2", "Chequing", "Yes", "Yes"),
    ("CIBC-JOINT-CHQ", "CIBC", "Joint Chequing", "Joint", "Chequing", "Yes", "Yes"),
    ("RBC-JOINT-SAV", "RBC", "Joint Goal Savings", "Joint", "Savings", "Yes", "Yes"),
    ("CIBC-P1-VISA", "CIBC", "Partner 1 Visa", "Partner 1", "Credit Card", "Yes", "Yes"),
    ("RBC-P2-VISA", "RBC", "Partner 2 Visa", "Partner 2", "Credit Card", "Yes", "Yes"),
    ("RBC-P1-TFSA", "RBC", "Partner 1 TFSA", "Partner 1", "Investment", "No", "Yes"),
    ("CIBC-P2-TFSA", "CIBC", "Partner 2 TFSA", "Partner 2", "Investment", "No", "Yes"),
]

RECONCILED_ACCOUNTS = [
    "CIBC-P1-CHQ",
    "RBC-P2-CHQ",
    "CIBC-JOINT-CHQ",
    "RBC-JOINT-SAV",
    "CIBC-P1-VISA",
    "RBC-P2-VISA",
]

BUDGETS = [
    ("Household", "Rent", 2200.00),
    ("Household", "Groceries", 700.00),
    ("Household", "Utilities", 300.00),
    ("Household", "Restaurants", 300.00),
    ("Household", "Transportation", 300.00),
    ("Household", "Subscriptions", 80.00),
    ("Household", "Household Supplies", 150.00),
    ("Partner 1", "Personal", 250.00),
    ("Partner 2", "Personal", 250.00),
]


def sample_transactions_and_reviews():
    transactions: list[tuple] = []
    reviews: list[list] = []

    def add(
        transaction_id: str,
        posted: date,
        account_id: str,
        merchant: str,
        amount: float,
        suggested_type: str,
        suggested_category: str,
        budget_for: str,
        *,
        status: str = "Verified",
        type_override: str = "",
        category_override: str = "",
        related: str = "",
        notes: str = "",
    ) -> None:
        transactions.append(
            (
                transaction_id,
                posted,
                account_id,
                merchant,
                amount,
                suggested_type,
                suggested_category,
            )
        )
        reviews.append(
            [
                transaction_id,
                status,
                type_override,
                category_override,
                budget_for,
                related,
                date(posted.year, posted.month, 1),
                notes,
            ]
        )

    # September: a completed and reconciled month.
    add("SEP-001", date(2026, 9, 1), "CIBC-P1-CHQ", "Partner 1 Payroll", 3500, "Income", "Employment", "Household")
    add("SEP-002", date(2026, 9, 1), "RBC-P2-CHQ", "Partner 2 Payroll", 3200, "Income", "Employment", "Household")
    add("SEP-003", date(2026, 9, 2), "CIBC-P1-CHQ", "Transfer to Joint", -2500, "Internal Transfer", "Transfer", "Household", related="P1-JOINT-SEP")
    add("SEP-004", date(2026, 9, 2), "CIBC-JOINT-CHQ", "Transfer from Partner 1", 2500, "Internal Transfer", "Transfer", "Household", related="P1-JOINT-SEP")
    add("SEP-005", date(2026, 9, 2), "RBC-P2-CHQ", "Transfer to Joint", -2200, "Internal Transfer", "Transfer", "Household", related="P2-JOINT-SEP")
    add("SEP-006", date(2026, 9, 2), "CIBC-JOINT-CHQ", "Transfer from Partner 2", 2200, "Internal Transfer", "Transfer", "Household", related="P2-JOINT-SEP")
    add("SEP-007", date(2026, 9, 3), "CIBC-JOINT-CHQ", "Rent Payment", -2200, "Expense", "Rent", "Household", related="RENT-2026-09")
    add("SEP-008", date(2026, 9, 5), "CIBC-JOINT-CHQ", "Utilities", -250, "Expense", "Utilities", "Household")
    add(
        "SEP-009",
        date(2026, 9, 6),
        "CIBC-JOINT-CHQ",
        "Transfer to Goal Savings",
        -500,
        "Internal Transfer",
        "Transfer",
        "Household",
        type_override="Sinking Fund Contribution",
        category_override="Emergency Fund",
        related="SINK-2026-09",
    )
    add("SEP-010", date(2026, 9, 6), "RBC-JOINT-SAV", "Transfer from Joint Chequing", 500, "Internal Transfer", "Transfer", "Household", related="SINK-2026-09")
    add(
        "SEP-011",
        date(2026, 9, 7),
        "CIBC-P1-CHQ",
        "TFSA Contribution",
        -500,
        "Internal Transfer",
        "Transfer",
        "Partner 1",
        type_override="Investment Contribution",
        category_override="TFSA",
        related="P1-INVEST-2026-09",
    )
    add(
        "SEP-012",
        date(2026, 9, 7),
        "RBC-P2-CHQ",
        "TFSA Contribution",
        -500,
        "Internal Transfer",
        "Transfer",
        "Partner 2",
        type_override="Investment Contribution",
        category_override="TFSA",
        related="P2-INVEST-2026-09",
    )
    add("SEP-013", date(2026, 9, 8), "CIBC-P1-VISA", "Grocery Store", -300, "Expense", "Groceries", "Household")
    add("SEP-014", date(2026, 9, 9), "RBC-P2-VISA", "Grocery Store", -300, "Expense", "Groceries", "Household")
    add("SEP-015", date(2026, 9, 12), "CIBC-P1-VISA", "Dinner with Friends", -250, "Expense", "Restaurants", "Household", related="DINNER-2026-09")
    add(
        "SEP-016",
        date(2026, 9, 13),
        "CIBC-P1-CHQ",
        "E-transfer from Friend",
        50,
        "Income",
        "Bank",
        "Household",
        type_override="Reimbursement Received",
        category_override="Restaurants",
        related="DINNER-2026-09",
    )
    add("SEP-017", date(2026, 9, 14), "CIBC-P1-VISA", "Transit and Fuel", -100, "Expense", "Transportation", "Household")
    add("SEP-018", date(2026, 9, 15), "RBC-P2-VISA", "Transit and Fuel", -100, "Expense", "Transportation", "Household")
    add("SEP-019", date(2026, 9, 16), "CIBC-JOINT-CHQ", "Streaming Services", -40, "Expense", "Subscriptions", "Household")
    add("SEP-020", date(2026, 9, 18), "CIBC-JOINT-CHQ", "Household Supplies", -100, "Expense", "Household Supplies", "Household")
    add("SEP-021", date(2026, 9, 20), "CIBC-P1-VISA", "Partner 1 Personal", -180, "Expense", "Personal", "Partner 1")
    add("SEP-022", date(2026, 9, 21), "RBC-P2-VISA", "Partner 2 Personal", -160, "Expense", "Personal", "Partner 2")
    add("SEP-023", date(2026, 9, 26), "CIBC-P1-CHQ", "Visa Payment", -830, "Internal Transfer", "Transfer", "Partner 1", related="P1-VISA-SEP")
    add("SEP-024", date(2026, 9, 26), "CIBC-P1-VISA", "Payment Received", 830, "Internal Transfer", "Transfer", "Partner 1", related="P1-VISA-SEP")
    add("SEP-025", date(2026, 9, 27), "RBC-P2-CHQ", "Visa Payment", -660, "Internal Transfer", "Transfer", "Partner 2", related="P2-VISA-SEP")
    add("SEP-026", date(2026, 9, 27), "RBC-P2-VISA", "Payment Received", 660, "Internal Transfer", "Transfer", "Partner 2", related="P2-VISA-SEP")

    # October: month-to-date with one transaction still awaiting review.
    add("OCT-001", date(2026, 10, 1), "CIBC-P1-CHQ", "Partner 1 Payroll", 3500, "Income", "Employment", "Household")
    add("OCT-002", date(2026, 10, 1), "RBC-P2-CHQ", "Partner 2 Payroll", 3200, "Income", "Employment", "Household")
    add("OCT-003", date(2026, 10, 2), "CIBC-P1-CHQ", "Transfer to Joint", -2500, "Internal Transfer", "Transfer", "Household", related="P1-JOINT-OCT")
    add("OCT-004", date(2026, 10, 2), "CIBC-JOINT-CHQ", "Transfer from Partner 1", 2500, "Internal Transfer", "Transfer", "Household", related="P1-JOINT-OCT")
    add("OCT-005", date(2026, 10, 2), "RBC-P2-CHQ", "Transfer to Joint", -2200, "Internal Transfer", "Transfer", "Household", related="P2-JOINT-OCT")
    add("OCT-006", date(2026, 10, 2), "CIBC-JOINT-CHQ", "Transfer from Partner 2", 2200, "Internal Transfer", "Transfer", "Household", related="P2-JOINT-OCT")
    add("OCT-007", date(2026, 10, 3), "CIBC-JOINT-CHQ", "Rent Payment", -2200, "Expense", "Rent", "Household", related="RENT-2026-10")
    add("OCT-008", date(2026, 10, 5), "CIBC-JOINT-CHQ", "Utilities", -180, "Expense", "Utilities", "Household")
    add(
        "OCT-009",
        date(2026, 10, 6),
        "CIBC-JOINT-CHQ",
        "Transfer to Goal Savings",
        -500,
        "Internal Transfer",
        "Transfer",
        "Household",
        type_override="Sinking Fund Contribution",
        category_override="Emergency Fund",
        related="SINK-2026-10",
    )
    add("OCT-010", date(2026, 10, 6), "RBC-JOINT-SAV", "Transfer from Joint Chequing", 500, "Internal Transfer", "Transfer", "Household", related="SINK-2026-10")
    add(
        "OCT-011",
        date(2026, 10, 7),
        "CIBC-P1-CHQ",
        "TFSA Contribution",
        -500,
        "Internal Transfer",
        "Transfer",
        "Partner 1",
        type_override="Investment Contribution",
        category_override="TFSA",
        related="P1-INVEST-2026-10",
    )
    add(
        "OCT-012",
        date(2026, 10, 7),
        "RBC-P2-CHQ",
        "TFSA Contribution",
        -500,
        "Internal Transfer",
        "Transfer",
        "Partner 2",
        type_override="Investment Contribution",
        category_override="TFSA",
        related="P2-INVEST-2026-10",
    )
    add("OCT-013", date(2026, 10, 8), "CIBC-P1-VISA", "Grocery Store", -250, "Expense", "Groceries", "Household")
    add("OCT-014", date(2026, 10, 8), "RBC-P2-VISA", "Grocery Store", -180, "Expense", "Groceries", "Household")
    add("OCT-015", date(2026, 10, 9), "CIBC-P1-VISA", "Dinner with Friends", -200, "Expense", "Restaurants", "Household", related="DINNER-2026-10")
    add(
        "OCT-016",
        date(2026, 10, 10),
        "CIBC-P1-CHQ",
        "E-transfer from Friend",
        100,
        "Income",
        "Bank",
        "Household",
        type_override="Reimbursement Received",
        category_override="Restaurants",
        related="DINNER-2026-10",
    )
    add("OCT-017", date(2026, 10, 10), "CIBC-P1-VISA", "Transit and Fuel", -80, "Expense", "Transportation", "Household")
    add("OCT-018", date(2026, 10, 10), "RBC-P2-VISA", "Transit and Fuel", -60, "Expense", "Transportation", "Household")
    add("OCT-019", date(2026, 10, 11), "CIBC-JOINT-CHQ", "Streaming Services", -40, "Expense", "Subscriptions", "Household")
    add("OCT-020", date(2026, 10, 11), "CIBC-JOINT-CHQ", "Household Supplies", -80, "Expense", "Household Supplies", "Household")
    add(
        "OCT-021",
        date(2026, 10, 12),
        "CIBC-P1-VISA",
        "Partner 1 Personal",
        -120,
        "Expense",
        "Personal",
        "Partner 1",
        status="Unreviewed",
        notes="Change this to Verified to close October.",
    )
    add("OCT-022", date(2026, 10, 12), "RBC-P2-VISA", "Partner 2 Personal", -90, "Expense", "Personal", "Partner 2")
    add("OCT-023", date(2026, 10, 13), "CIBC-P1-CHQ", "Visa Payment", -600, "Internal Transfer", "Transfer", "Partner 1", related="P1-VISA-OCT")
    add("OCT-024", date(2026, 10, 13), "CIBC-P1-VISA", "Payment Received", 600, "Internal Transfer", "Transfer", "Partner 1", related="P1-VISA-OCT")
    add("OCT-025", date(2026, 10, 13), "RBC-P2-CHQ", "Visa Payment", -500, "Internal Transfer", "Transfer", "Partner 2", related="P2-VISA-OCT")
    add("OCT-026", date(2026, 10, 13), "RBC-P2-VISA", "Payment Received", 500, "Internal Transfer", "Transfer", "Partner 2", related="P2-VISA-OCT")

    return transactions, reviews


def account_activity(transactions, month: int) -> dict[str, float]:
    totals: defaultdict[str, float] = defaultdict(float)
    for transaction in transactions:
        _, posted, account_id, _, amount, _, _ = transaction
        if posted.month == month:
            totals[account_id] += amount
    return dict(totals)


def build_balances(transactions):
    september_opening = {
        "CIBC-P1-CHQ": 2800.00,
        "RBC-P2-CHQ": 2500.00,
        "CIBC-JOINT-CHQ": 1200.00,
        "RBC-JOINT-SAV": 4500.00,
        "CIBC-P1-VISA": -300.00,
        "RBC-P2-VISA": -250.00,
        "RBC-P1-TFSA": 14300.00,
        "CIBC-P2-TFSA": 11300.00,
    }
    september_activity = account_activity(transactions, 9)
    september_ending = {
        account_id: opening + september_activity.get(account_id, 0)
        for account_id, opening in september_opening.items()
    }
    september_ending["RBC-P1-TFSA"] = 14800.00
    september_ending["CIBC-P2-TFSA"] = 11850.00

    october_activity = account_activity(transactions, 10)
    october_ending = {
        account_id: opening + october_activity.get(account_id, 0)
        for account_id, opening in september_ending.items()
    }
    october_ending["RBC-P1-TFSA"] = 15450.00
    october_ending["CIBC-P2-TFSA"] = 12390.00

    rows = []
    for snapshot_date, balances, source in [
        (date(2026, 9, 1), september_opening, "Opening statement"),
        (date(2026, 9, 30), september_ending, "Month-end statement"),
        (date(2026, 10, 13), october_ending, "Online balance / market value"),
    ]:
        for account_id, balance in balances.items():
            note = "Market value" if account_id.endswith("TFSA") else ""
            rows.append([snapshot_date, account_id, balance, source, note])
    return rows, september_opening, september_ending, october_ending


def make_lists(wb: Workbook) -> None:
    ws = wb.create_sheet("Lists")
    lists = [
        ("Statuses", ["Verified", "Unreviewed", "Needs Correction", "Duplicate", "Excluded"]),
        (
            "Types",
            [
                "Income",
                "Expense",
                "Reimbursement Received",
                "Sinking Fund Contribution",
                "Investment Contribution",
                "Internal Transfer",
            ],
        ),
        (
            "Categories",
            [
                "Employment",
                "Rent",
                "Groceries",
                "Utilities",
                "Restaurants",
                "Transportation",
                "Subscriptions",
                "Household Supplies",
                "Personal",
                "Emergency Fund",
                "TFSA",
                "Transfer",
                "Bank",
            ],
        ),
        ("BudgetFor", ["Household", "Partner 1", "Partner 2"]),
    ]
    for column, (header, values) in enumerate(lists, start=1):
        ws.cell(1, column, header)
        for row, value in enumerate(values, start=2):
            ws.cell(row, column, value)

    wb.defined_names.add(DefinedName("CoupleStatusList", attr_text="'Lists'!$A$2:$A$6"))
    wb.defined_names.add(DefinedName("CoupleTypeList", attr_text="'Lists'!$B$2:$B$7"))
    wb.defined_names.add(DefinedName("CoupleCategoryList", attr_text="'Lists'!$C$2:$C$14"))
    wb.defined_names.add(DefinedName("BudgetForList", attr_text="'Lists'!$D$2:$D$4"))
    ws.sheet_state = "hidden"


def make_start_here(wb: Workbook) -> None:
    ws = wb.create_sheet("Start Here")
    ws.sheet_view.showGridLines = False
    for column, width in {"A": 4, "B": 28, "C": 28, "D": 28, "E": 28, "F": 4}.items():
        ws.column_dimensions[column].width = width

    ws.merge_cells("B2:E2")
    ws["B2"] = "Couple Monthly Budget — Sample Workbook"
    ws["B2"].fill = PatternFill("solid", fgColor=NAVY)
    ws["B2"].font = Font(color=WHITE, size=19, bold=True)
    ws["B2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 34
    ws.merge_cells("B3:E3")
    ws["B3"] = (
        "One household workbook for two people, multiple institutions, "
        "shared expenses, personal spending, and investments."
    )
    ws["B3"].font = Font(color=GRAY, italic=True)
    ws["B3"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[3].height = 32

    section_header(ws, "B6:E6", "The workbook has three connected layers")
    layers = [
        ("1. Bank facts", "Transactions contain posted dates, accounts, merchants, and signed amounts."),
        ("2. Your decisions", "Review Log stores verification, overrides, who the spending was for, and related IDs."),
        ("3. Monthly views", "September and October summarize the final reviewed transactions."),
    ]
    for row, (label, detail) in enumerate(layers, start=7):
        ws[f"B{row}"] = label
        ws[f"B{row}"].font = Font(bold=True, color=NAVY)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        ws.cell(row, 3, detail)
        ws.cell(row, 3).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 28

    section_header(ws, "B12:E12", "What changes for a couple")
    couple_rules = [
        ("Account Owner", "Who legally owns or uses the bank or investment account."),
        ("Budget For", "Household, Partner 1, or Partner 2—the person or group benefiting from the spending."),
        ("Internal Transfer", "Money moved between included family accounts; never counted as family income or spending."),
        ("Balance Snapshot", "A dated bank balance or investment market value used for current position and history."),
    ]
    for row, (field, explanation) in enumerate(couple_rules, start=13):
        ws[f"B{row}"] = field
        ws[f"B{row}"].font = Font(bold=True)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        ws.cell(row, 3, explanation)
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
        ws.cell(row, 3).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 28

    section_header(ws, "B19:E19", "How to try the sample")
    steps = [
        "Open October. Yellow cells are inputs; blue cells calculate automatically.",
        "Open Accounts to see current CIBC, RBC, credit-card, and investment values.",
        "Open Review Log and change OCT-021 from Unreviewed to Verified.",
        "Return to October. With all statement differences at $0, the month becomes Reconciled.",
        "Replace the sample names Partner 1 and Partner 2 with your names when adapting the workbook.",
    ]
    for row, text in enumerate(steps, start=20):
        ws[f"B{row}"] = f"{row - 19}."
        ws[f"B{row}"].font = Font(bold=True, color=NAVY)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        ws.cell(row, 3, text)
        ws.cell(row, 3).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 28

    section_header(ws, "B27:E27", "Open a sheet")
    for cell, label, target in [
        ("B28", "October", "October"),
        ("C28", "Accounts", "Accounts"),
        ("D28", "Transactions", "Transactions"),
        ("E28", "Review Log", "Review Log"),
    ]:
        add_nav_link(ws, cell, label, target)
        ws[cell].alignment = Alignment(horizontal="center")
    ws.sheet_properties.tabColor = GREEN


def make_balances_sheet(wb: Workbook, balance_rows) -> None:
    ws = wb.create_sheet("Balance Snapshots")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:E1")
    ws["A1"] = "Balance Snapshots"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:E2")
    ws["A2"] = (
        "Enter one dated balance per account. Investment balances are market values, "
        "which cannot be calculated from transaction CSVs alone."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=YELLOW)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 32
    add_nav_link(ws, "A3", "← Start Here", "Start Here")
    add_nav_link(ws, "B3", "Accounts", "Accounts")
    add_nav_link(ws, "C3", "October", "October")

    headers = ["As Of Date", "Account ID", "Balance", "Source", "Notes"]
    for column, header in enumerate(headers, start=1):
        ws.cell(4, column, header)
    for row_number, values in enumerate(balance_rows, start=5):
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
            ws.cell(row_number, column).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row_number, 1).number_format = DATE
        ws.cell(row_number, 3).number_format = MONEY

    last_row = 4 + len(balance_rows)
    add_table(ws, "tblBalances", f"A4:E{last_row}")
    for column, width in enumerate([16, 20, 18, 30, 26], start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.sheet_properties.tabColor = ORANGE


def make_accounts_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Accounts")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A10"
    ws.merge_cells("A1:I1")
    ws["A1"] = "Family Accounts & Current Values"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:I2")
    ws["A2"] = (
        "Account details identify the institution and owner. Latest values come from Balance Snapshots—not from assuming a CSV contains the balance."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
    ws["A2"].alignment = Alignment(wrap_text=True)
    add_nav_link(ws, "A3", "← Start Here", "Start Here")
    add_nav_link(ws, "B3", "Balance Snapshots", "Balance Snapshots")
    add_nav_link(ws, "C3", "October", "October")

    summary = [
        ("A5:B5", "Cash"),
        ("C5:D5", "Investments"),
        ("E5:F5", "Credit-Card Debt"),
        ("G5:I5", "Family Net Worth"),
    ]
    for cell_range, label in summary:
        ws.merge_cells(cell_range)
        cell = ws[cell_range.split(":")[0]]
        cell.value = label
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center")
    for cell_range in ["A6:B7", "C6:D7", "E6:F7", "G6:I7"]:
        ws.merge_cells(cell_range)
        cell = ws[cell_range.split(":")[0]]
        cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
        cell.font = Font(color=NAVY, bold=True, size=17)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.number_format = MONEY
    ws["A6"] = '=SUMIFS(tblAccounts[Latest Balance],tblAccounts[Type],"Chequing")+SUMIFS(tblAccounts[Latest Balance],tblAccounts[Type],"Savings")'
    ws["C6"] = '=SUMIFS(tblAccounts[Latest Balance],tblAccounts[Type],"Investment")'
    ws["E6"] = '=-SUMIFS(tblAccounts[Latest Balance],tblAccounts[Type],"Credit Card")'
    ws["G6"] = '=SUMIFS(tblAccounts[Latest Balance],tblAccounts[Include in Net Worth],"Yes")'

    headers = [
        "Account ID",
        "Institution",
        "Account Name",
        "Owner",
        "Type",
        "Reconcile Monthly",
        "Include in Net Worth",
        "Latest Balance",
        "Balance As Of",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(9, column, header)
    for row_number, account in enumerate(ACCOUNTS, start=10):
        for column, value in enumerate(account, start=1):
            ws.cell(row_number, column, value)
            ws.cell(row_number, column).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(
            row_number,
            8,
            f'=IFERROR(SUMIFS(tblBalances[Balance],tblBalances[Account ID],A{row_number},'
            f'tblBalances[As Of Date],MAXIFS(tblBalances[As Of Date],'
            f'tblBalances[Account ID],A{row_number})),"")',
        )
        ws.cell(
            row_number,
            9,
            f'=IFERROR(MAXIFS(tblBalances[As Of Date],tblBalances[Account ID],A{row_number}),"")',
        )
        ws.cell(row_number, 8).fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws.cell(row_number, 9).fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws.cell(row_number, 8).number_format = MONEY
        ws.cell(row_number, 9).number_format = DATE

    last_row = 9 + len(ACCOUNTS)
    add_table(ws, "tblAccounts", f"A9:I{last_row}")
    for column, width in enumerate([20, 14, 25, 15, 16, 20, 21, 18, 17], start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.sheet_properties.tabColor = GREEN


def make_review_sheet(wb: Workbook, review_rows) -> None:
    ws = wb.create_sheet("Review Log")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:H1")
    ws["A1"] = "Transaction Review & Household Allocation"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:H2")
    ws["A2"] = (
        "Use Budget For to distinguish shared household spending from each partner's personal spending. "
        "Manual corrections remain attached through Transaction ID."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=YELLOW)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 34
    add_nav_link(ws, "A3", "← Start Here", "Start Here")
    add_nav_link(ws, "B3", "October", "October")
    add_nav_link(ws, "C3", "Transactions", "Transactions")

    headers = [
        "Transaction ID",
        "Status",
        "Type Override",
        "Category Override",
        "Budget For",
        "Related Transaction",
        "Applies To Month",
        "Notes",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(4, column, header)
    for row_number, values in enumerate(review_rows, start=5):
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
            ws.cell(row_number, column).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row_number, 7).number_format = MONTH

    last_row = 4 + len(review_rows)
    add_table(ws, "tblCoupleReview", f"A4:H{last_row}")
    for column, width in enumerate([16, 20, 30, 24, 18, 25, 20, 38], start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    validations = [
        (DataValidation(type="list", formula1="=CoupleStatusList", allow_blank=False), f"B5:B{last_row}"),
        (DataValidation(type="list", formula1="=CoupleTypeList", allow_blank=True), f"C5:C{last_row}"),
        (DataValidation(type="list", formula1="=CoupleCategoryList", allow_blank=True), f"D5:D{last_row}"),
        (DataValidation(type="list", formula1="=BudgetForList", allow_blank=False), f"E5:E{last_row}"),
    ]
    for validation, cell_range in validations:
        ws.add_data_validation(validation)
        validation.add(cell_range)
    ws.conditional_formatting.add(
        f"B5:B{last_row}",
        FormulaRule(formula=['B5="Verified"'], fill=PatternFill("solid", fgColor=PALE_GREEN)),
    )
    ws.conditional_formatting.add(
        f"B5:B{last_row}",
        FormulaRule(formula=['B5<>"Verified"'], fill=PatternFill("solid", fgColor=PALE_RED)),
    )
    ws.sheet_properties.tabColor = ORANGE


def make_transactions_sheet(wb: Workbook, transactions) -> None:
    ws = wb.create_sheet("Transactions")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:P1")
    ws["A1"] = "Final Household Transactions"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:P2")
    ws["A2"] = (
        "A:G simulate cleaned Power Query output. H:P combine account ownership with the Review Log. "
        "Monthly pages use the final calculated columns."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws["A2"].alignment = Alignment(wrap_text=True)
    add_nav_link(ws, "A3", "← Start Here", "Start Here")
    add_nav_link(ws, "B3", "October", "October")
    add_nav_link(ws, "C3", "Review Log", "Review Log")

    headers = [
        "Transaction ID",
        "Posted Date",
        "Account ID",
        "Merchant",
        "Amount",
        "Suggested Type",
        "Suggested Category",
        "Institution",
        "Account Owner",
        "Final Type",
        "Final Category",
        "Budget For",
        "Related Transaction",
        "Applies To Month",
        "Review Status",
        "Notes",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(4, column, header)

    for row_number, transaction in enumerate(transactions, start=5):
        for column, value in enumerate(transaction, start=1):
            ws.cell(row_number, column, value)
        ws.cell(row_number, 2).number_format = DATE
        ws.cell(row_number, 5).number_format = MONEY
        ws.cell(row_number, 8, f'=IFERROR(INDEX(tblAccounts[Institution],MATCH(C{row_number},tblAccounts[Account ID],0)),"")')
        ws.cell(row_number, 9, f'=IFERROR(INDEX(tblAccounts[Owner],MATCH(C{row_number},tblAccounts[Account ID],0)),"")')
        ws.cell(
            row_number,
            10,
            f'=IFERROR(IF(INDEX(tblCoupleReview[Type Override],MATCH(A{row_number},'
            f'tblCoupleReview[Transaction ID],0))="",F{row_number},'
            f'INDEX(tblCoupleReview[Type Override],MATCH(A{row_number},'
            f'tblCoupleReview[Transaction ID],0))),F{row_number})',
        )
        ws.cell(
            row_number,
            11,
            f'=IFERROR(IF(INDEX(tblCoupleReview[Category Override],MATCH(A{row_number},'
            f'tblCoupleReview[Transaction ID],0))="",G{row_number},'
            f'INDEX(tblCoupleReview[Category Override],MATCH(A{row_number},'
            f'tblCoupleReview[Transaction ID],0))),G{row_number})',
        )
        ws.cell(row_number, 12, f'=IFERROR(INDEX(tblCoupleReview[Budget For],MATCH(A{row_number},tblCoupleReview[Transaction ID],0)),"Household")')
        ws.cell(row_number, 13, f'=IFERROR(INDEX(tblCoupleReview[Related Transaction],MATCH(A{row_number},tblCoupleReview[Transaction ID],0)),"")')
        ws.cell(
            row_number,
            14,
            f'=IFERROR(IF(INDEX(tblCoupleReview[Applies To Month],MATCH(A{row_number},'
            f'tblCoupleReview[Transaction ID],0))=0,DATE(YEAR(B{row_number}),MONTH(B{row_number}),1),'
            f'INDEX(tblCoupleReview[Applies To Month],MATCH(A{row_number},'
            f'tblCoupleReview[Transaction ID],0))),DATE(YEAR(B{row_number}),MONTH(B{row_number}),1))',
        )
        ws.cell(row_number, 14).number_format = MONTH
        ws.cell(row_number, 15, f'=IFERROR(INDEX(tblCoupleReview[Status],MATCH(A{row_number},tblCoupleReview[Transaction ID],0)),"Unreviewed")')
        ws.cell(row_number, 16, f'=IFERROR(INDEX(tblCoupleReview[Notes],MATCH(A{row_number},tblCoupleReview[Transaction ID],0)),"")')
        for column in range(8, 17):
            ws.cell(row_number, column).fill = PatternFill("solid", fgColor=PALE_BLUE)

    last_row = 4 + len(transactions)
    add_table(ws, "tblCoupleTransactions", f"A4:P{last_row}")
    widths = [16, 15, 20, 30, 15, 28, 23, 14, 17, 30, 23, 17, 25, 20, 19, 38]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.conditional_formatting.add(
        f"O5:O{last_row}",
        FormulaRule(formula=['O5="Verified"'], fill=PatternFill("solid", fgColor=PALE_GREEN)),
    )
    ws.conditional_formatting.add(
        f"O5:O{last_row}",
        FormulaRule(formula=['O5<>"Verified"'], fill=PatternFill("solid", fgColor=PALE_RED)),
    )
    ws.sheet_properties.tabColor = GRAY


def make_month_sheet(
    wb: Workbook,
    sheet_name: str,
    month_start: date,
    opening_balances: dict[str, float],
    ending_balances: dict[str, float],
    expected_income: float,
    sinking_goal: float,
    investment_goal: float,
) -> None:
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A10"
    ws.sheet_properties.tabColor = BLUE
    ws["M1"] = month_start
    ws["M1"].number_format = MONTH
    ws.column_dimensions["M"].hidden = True
    widths = {
        "A": 17,
        "B": 18,
        "C": 15,
        "D": 16,
        "E": 15,
        "F": 15,
        "G": 15,
        "H": 12,
        "I": 3,
        "J": 23,
        "K": 20,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.merge_cells("A1:K1")
    ws["A1"] = f"{sheet_name} 2026 — Couple Budget Check"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells("A2:K2")
    ws["A2"] = (
        "Family totals include both partners. Budget For separates shared household spending "
        "from each partner's personal spending."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
    ws["A2"].font = Font(color=GRAY, italic=True)
    ws["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    add_nav_link(ws, "A3", "← Start Here", "Start Here")
    add_nav_link(ws, "B3", "Accounts", "Accounts")
    add_nav_link(ws, "C3", "Transactions", "Transactions")
    add_nav_link(ws, "D3", "Review Log", "Review Log")
    other_month = "September" if sheet_name == "October" else "October"
    add_nav_link(ws, "E3", other_month, other_month)
    ws.merge_cells("G3:K3")
    ws["G3"] = "Yellow = editable   |   Blue = calculated"
    ws["G3"].font = Font(color=GRAY, italic=True)
    ws["G3"].alignment = Alignment(horizontal="right")

    cards = [
        ("A5:B5", "A6:B7", "A8:B8", "Income Received"),
        ("C5:D5", "C6:D7", "C8:D8", "Net Spending"),
        ("E5:F5", "E6:F7", "E8:F8", "Budget Remaining"),
        ("G5:H5", "G6:H7", "G8:H8", "Goals Funded"),
        ("J5:K5", "J6:K7", "J8:K8", "Unallocated Income"),
    ]
    for header_range, value_range, detail_range, label in cards:
        ws.merge_cells(header_range)
        ws.merge_cells(value_range)
        ws.merge_cells(detail_range)
        header = ws[header_range.split(":")[0]]
        value = ws[value_range.split(":")[0]]
        detail = ws[detail_range.split(":")[0]]
        header.value = label
        header.fill = PatternFill("solid", fgColor=BLUE)
        header.font = Font(color=WHITE, bold=True)
        header.alignment = Alignment(horizontal="center")
        value.fill = PatternFill("solid", fgColor=PALE_BLUE)
        value.font = Font(color=NAVY, bold=True, size=17)
        value.alignment = Alignment(horizontal="center", vertical="center")
        value.number_format = MONEY_POSITIVE
        detail.fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
        detail.font = Font(color=GRAY, italic=True, size=9)
        detail.alignment = Alignment(horizontal="center")

    ws["A6"] = "=K24"
    ws["C6"] = "=K27"
    ws["E6"] = "=SUM(G12:G20)"
    ws["G6"] = "=K28+K29"
    ws["J6"] = "=K30"
    ws["A8"] = '="of "&TEXT(K11,"$#,##0")&" expected"'
    ws["C8"] = '="Gross "&TEXT(K25,"$#,##0")&" less "&TEXT(K26,"$#,##0")'
    ws["E8"] = '="of "&TEXT(SUM(C12:C20),"$#,##0")&" planned"'
    ws["G8"] = '="Goal "&TEXT(K12+K13,"$#,##0")'
    ws["J8"] = "After spending and goals"

    section_header(ws, "A10:H10", "Budget by household or person — edit yellow Budget cells")
    headers = ["Budget For", "Category", "Budget", "Gross Spend", "Reimbursed", "Net Spend", "Remaining", "% Used"]
    for column, header in enumerate(headers, start=1):
        ws.cell(11, column, header)
        ws.cell(11, column).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(11, column).font = Font(color=WHITE, bold=True)
        ws.cell(11, column).alignment = Alignment(horizontal="center", wrap_text=True)
    for row, (budget_for, category, budget) in enumerate(BUDGETS, start=12):
        ws.cell(row, 1, budget_for)
        ws.cell(row, 2, category)
        ws.cell(row, 3, budget)
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(
            row,
            4,
            f'=-SUMIFS(tblCoupleTransactions[Amount],tblCoupleTransactions[Final Type],"Expense",'
            f'tblCoupleTransactions[Final Category],$B{row},tblCoupleTransactions[Budget For],$A{row},'
            f'tblCoupleTransactions[Applies To Month],$M$1)',
        )
        ws.cell(
            row,
            5,
            f'=SUMIFS(tblCoupleTransactions[Amount],tblCoupleTransactions[Final Type],'
            f'"Reimbursement Received",tblCoupleTransactions[Final Category],$B{row},'
            f'tblCoupleTransactions[Budget For],$A{row},tblCoupleTransactions[Applies To Month],$M$1)',
        )
        ws.cell(row, 6, f"=D{row}-E{row}")
        ws.cell(row, 7, f"=C{row}-F{row}")
        ws.cell(row, 8, f'=IFERROR(F{row}/C{row},0)')
        for column in range(3, 8):
            ws.cell(row, column).number_format = MONEY_POSITIVE
        ws.cell(row, 8).number_format = PERCENT
        for column in range(4, 9):
            ws.cell(row, column).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
    total_row = 21
    ws.cell(total_row, 1, "TOTAL")
    ws.cell(total_row, 1).font = Font(bold=True)
    for column in range(3, 8):
        letter = get_column_letter(column)
        ws.cell(total_row, column, f"=SUM({letter}12:{letter}20)")
        ws.cell(total_row, column).number_format = MONEY_POSITIVE
        ws.cell(total_row, column).font = Font(bold=True)
        ws.cell(total_row, column).fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws.cell(total_row, 8, "=IFERROR(F21/C21,0)")
    ws.cell(total_row, 8).number_format = PERCENT
    ws.cell(total_row, 8).font = Font(bold=True)
    ws.cell(total_row, 8).fill = PatternFill("solid", fgColor=PALE_BLUE)
    set_all_borders(ws, "A11:H21")
    ws.conditional_formatting.add(
        "G12:G21",
        CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=PALE_RED), font=Font(color=RED)),
    )
    ws.conditional_formatting.add(
        "G12:G21",
        CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=PatternFill("solid", fgColor=PALE_GREEN)),
    )
    ws.conditional_formatting.add(
        "H12:H21",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=BLUE),
    )

    section_header(ws, "J10:K10", "Month setup & verification")
    setup_labels = [
        "Expected income",
        "Sinking-fund goal",
        "Investment goal",
        "Transactions through",
        "Transactions",
        "Verified",
        "Needs review",
        "Statement difference",
        "Month status",
    ]
    for row, label in enumerate(setup_labels, start=11):
        ws.cell(row, 10, label)
        ws.cell(row, 10).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
        ws.cell(row, 11).fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws.cell(row, 11).alignment = Alignment(horizontal="center")
    ws["K11"] = expected_income
    ws["K12"] = sinking_goal
    ws["K13"] = investment_goal
    for cell in ["K11", "K12", "K13"]:
        ws[cell].fill = PatternFill("solid", fgColor=YELLOW)
        ws[cell].number_format = MONEY_POSITIVE
        ws[cell].comment = Comment("Editable monthly plan.", "Cursor")
    ws["K14"] = '=IFERROR(AGGREGATE(14,6,tblCoupleTransactions[Posted Date]/(tblCoupleTransactions[Applies To Month]=$M$1),1),"")'
    ws["K14"].number_format = DATE
    ws["K15"] = '=COUNTIFS(tblCoupleTransactions[Applies To Month],$M$1,tblCoupleTransactions[Review Status],"<>Excluded")'
    ws["K16"] = '=COUNTIFS(tblCoupleTransactions[Applies To Month],$M$1,tblCoupleTransactions[Review Status],"Verified")'
    ws["K17"] = "=K15-K16"
    ws["K18"] = "=SUMPRODUCT(ABS(G25:G30))"
    ws["K18"].number_format = MONEY_POSITIVE
    ws["K19"] = '=IF(AND(K17=0,K18<0.01),"RECONCILED","NEEDS REVIEW")'
    ws["K19"].font = Font(bold=True)
    set_all_borders(ws, "J10:K19")
    ws.conditional_formatting.add(
        "K19",
        FormulaRule(formula=['K19="RECONCILED"'], fill=PatternFill("solid", fgColor=PALE_GREEN), font=Font(color=GREEN, bold=True)),
    )
    ws.conditional_formatting.add(
        "K19",
        FormulaRule(formula=['K19<>"RECONCILED"'], fill=PatternFill("solid", fgColor=PALE_RED), font=Font(color=RED, bold=True)),
    )

    section_header(ws, "A23:H23", "Bank statement check — internal transfers count here, but not in spending")
    reconcile_headers = ["Account ID", "Owner", "Opening", "Imported Activity", "Calculated Ending", "Statement Ending", "Difference", "Status"]
    for column, header in enumerate(reconcile_headers, start=1):
        ws.cell(24, column, header)
        ws.cell(24, column).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(24, column).font = Font(color=WHITE, bold=True)
        ws.cell(24, column).alignment = Alignment(horizontal="center", wrap_text=True)
    account_owner = {account[0]: account[3] for account in ACCOUNTS}
    for row, account_id in enumerate(RECONCILED_ACCOUNTS, start=25):
        ws.cell(row, 1, account_id)
        ws.cell(row, 2, account_owner[account_id])
        ws.cell(row, 3, opening_balances[account_id])
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(
            row,
            4,
            f'=SUMIFS(tblCoupleTransactions[Amount],tblCoupleTransactions[Account ID],$A{row},'
            f'tblCoupleTransactions[Posted Date],">="&$M$1,tblCoupleTransactions[Posted Date],"<"&EDATE($M$1,1))',
        )
        ws.cell(row, 5, f"=C{row}+D{row}")
        ws.cell(row, 6, ending_balances[account_id])
        ws.cell(row, 6).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row, 7, f"=F{row}-E{row}")
        ws.cell(row, 8, f'=IF(ABS(G{row})<0.01,"Matched","Check")')
        for column in range(3, 8):
            ws.cell(row, column).number_format = MONEY
        for column in [4, 5, 7, 8]:
            ws.cell(row, column).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
    set_all_borders(ws, "A24:H30")
    ws.conditional_formatting.add(
        "G25:G30",
        FormulaRule(formula=["ABS(G25)<0.01"], fill=PatternFill("solid", fgColor=PALE_GREEN)),
    )
    ws.conditional_formatting.add(
        "G25:G30",
        FormulaRule(formula=["ABS(G25)>=0.01"], fill=PatternFill("solid", fgColor=PALE_RED)),
    )

    section_header(ws, "J23:K23", "How family income was used")
    labels = [
        "Income received",
        "Gross spending",
        "Reimbursements",
        "Net spending",
        "Sinking funds",
        "Investments",
        "Unallocated income",
    ]
    formulas = [
        '=SUMIFS(tblCoupleTransactions[Amount],tblCoupleTransactions[Final Type],"Income",tblCoupleTransactions[Applies To Month],$M$1)',
        '=-SUMIFS(tblCoupleTransactions[Amount],tblCoupleTransactions[Final Type],"Expense",tblCoupleTransactions[Applies To Month],$M$1)',
        '=SUMIFS(tblCoupleTransactions[Amount],tblCoupleTransactions[Final Type],"Reimbursement Received",tblCoupleTransactions[Applies To Month],$M$1)',
        "=K25-K26",
        '=-SUMIFS(tblCoupleTransactions[Amount],tblCoupleTransactions[Final Type],"Sinking Fund Contribution",tblCoupleTransactions[Applies To Month],$M$1)',
        '=-SUMIFS(tblCoupleTransactions[Amount],tblCoupleTransactions[Final Type],"Investment Contribution",tblCoupleTransactions[Applies To Month],$M$1)',
        "=K24-K27-K28-K29",
    ]
    for row, (label, formula) in enumerate(zip(labels, formulas), start=24):
        ws.cell(row, 10, label)
        ws.cell(row, 11, formula)
        ws.cell(row, 11).number_format = MONEY_POSITIVE
        ws.cell(row, 10).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
        ws.cell(row, 11).fill = PatternFill("solid", fgColor=PALE_GREEN if row == 30 else PALE_BLUE)
        if row == 30:
            ws.cell(row, 10).font = Font(bold=True)
            ws.cell(row, 11).font = Font(bold=True)
    set_all_borders(ws, "J23:K30")

    section_header(ws, "A33:E33", "Who the spending was for")
    for column, header in enumerate(["Budget For", "Budget", "Net Spend", "Remaining", "Meaning"], start=1):
        ws.cell(34, column, header)
        ws.cell(34, column).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(34, column).font = Font(color=WHITE, bold=True)
    meanings = {
        "Household": "Shared living costs",
        "Partner 1": "Partner 1 personal spending",
        "Partner 2": "Partner 2 personal spending",
    }
    for row, budget_for in enumerate(["Household", "Partner 1", "Partner 2"], start=35):
        ws.cell(row, 1, budget_for)
        ws.cell(row, 2, f'=SUMIF($A$12:$A$20,A{row},$C$12:$C$20)')
        ws.cell(row, 3, f'=SUMIF($A$12:$A$20,A{row},$F$12:$F$20)')
        ws.cell(row, 4, f"=B{row}-C{row}")
        ws.cell(row, 5, meanings[budget_for])
        for column in range(2, 5):
            ws.cell(row, column).number_format = MONEY_POSITIVE
    set_all_borders(ws, "A34:E37")

    section_header(ws, "J33:K33", "Current family position")
    current_labels = ["Cash", "Investments", "Credit-card debt", "Net worth", "Balances as of"]
    current_formulas = [
        '=SUMIFS(tblAccounts[Latest Balance],tblAccounts[Type],"Chequing")+SUMIFS(tblAccounts[Latest Balance],tblAccounts[Type],"Savings")',
        '=SUMIFS(tblAccounts[Latest Balance],tblAccounts[Type],"Investment")',
        '=-SUMIFS(tblAccounts[Latest Balance],tblAccounts[Type],"Credit Card")',
        '=SUMIFS(tblAccounts[Latest Balance],tblAccounts[Include in Net Worth],"Yes")',
        '=MAX(tblAccounts[Balance As Of])',
    ]
    for row, (label, formula) in enumerate(zip(current_labels, current_formulas), start=34):
        ws.cell(row, 10, label)
        ws.cell(row, 11, formula)
        ws.cell(row, 10).fill = PatternFill("solid", fgColor=VERY_PALE_BLUE)
        ws.cell(row, 11).fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws.cell(row, 11).number_format = DATE if row == 38 else MONEY
    set_all_borders(ws, "J33:K38")

    ws.merge_cells("A40:K40")
    ws["A40"] = (
        "A transfer between Partner 1, Partner 2, joint accounts, credit cards, or savings changes account balances—but not family spending."
    )
    ws["A40"].fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    ws["A40"].font = Font(color=GRAY, italic=True)
    ws["A40"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[40].height = 28
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:K40"


def build_workbook() -> None:
    transactions, reviews = sample_transactions_and_reviews()
    balance_rows, september_opening, september_ending, october_ending = build_balances(transactions)

    wb = Workbook()
    wb.remove(wb.active)
    make_lists(wb)
    make_start_here(wb)
    make_balances_sheet(wb, balance_rows)
    make_accounts_sheet(wb)
    make_review_sheet(wb, reviews)
    make_transactions_sheet(wb, transactions)
    make_month_sheet(
        wb,
        "September",
        date(2026, 9, 1),
        september_opening,
        september_ending,
        expected_income=6700.00,
        sinking_goal=500.00,
        investment_goal=1000.00,
    )
    make_month_sheet(
        wb,
        "October",
        date(2026, 10, 1),
        september_ending,
        october_ending,
        expected_income=6700.00,
        sinking_goal=500.00,
        investment_goal=1000.00,
    )

    order = [
        "Start Here",
        "October",
        "September",
        "Accounts",
        "Balance Snapshots",
        "Transactions",
        "Review Log",
        "Lists",
    ]
    wb._sheets = [wb[name] for name in order]
    wb.active = wb.sheetnames.index("October")
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == [
        "Start Here",
        "October",
        "September",
        "Accounts",
        "Balance Snapshots",
        "Transactions",
        "Review Log",
        "Lists",
    ]
    assert wb["Lists"].sheet_state == "hidden"
    assert "tblAccounts" in wb["Accounts"].tables
    assert "tblBalances" in wb["Balance Snapshots"].tables
    assert "tblCoupleTransactions" in wb["Transactions"].tables
    assert "tblCoupleReview" in wb["Review Log"].tables
    assert wb["October"]["K19"].value.startswith("=IF(")
    assert wb["October"]["C12"].fill.fgColor.rgb.endswith(YELLOW)
    assert wb["October"]["G25"].value == "=F25-E25"
    review_statuses = [
        cell.value for cell in wb["Review Log"]["B"] if cell.row >= 5
    ]
    assert review_statuses.count("Unreviewed") == 1


if __name__ == "__main__":
    build_workbook()
    print(f"Created {OUTPUT}")
